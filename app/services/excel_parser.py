import os
import sys
import pandas as pd
from app.models.work_report_model import WorkReportEntry

# 可選依賴：openpyxl (用於 XLSX 讀取)
OPENPYXL_AVAILABLE = False
try:
    from openpyxl import load_workbook
    OPENPYXL_AVAILABLE = True
except ImportError:
    pass

# 可選依賴：xlrd (用於 XLS 讀取)
XLRD_AVAILABLE = False
try:
    import xlrd
    XLRD_AVAILABLE = True
except ImportError:
    pass

# 可選依賴：pywin32 (用於 XLS COM 接口)
WIN32_AVAILABLE = False
try:
    import win32com.client as win32
    import pythoncom
    WIN32_AVAILABLE = True
except ImportError:
    pass

class ExcelParser:
    """Excel 結構化解析器 - 專門用於工作報告的資料提取和分組"""
    
    # 可能的填單人員欄位名稱（關鍵字匹配）
    USER_FIELD_KEYWORDS = ['填單人員', '填單人', '需求人', '負責人', '人員', '姓名']
    
    @staticmethod
    def parse_with_pandas(file_path):
        """使用 Pandas 解析 Excel"""
        try:
            print(f"[INFO] 使用 Pandas 解析: {file_path}")
            
            # 1. 讀取檔案，不預設 header，以便尋找正確的標題行
            # 支援 xls and xlsx
            df_raw = pd.read_excel(file_path, header=None)
            
            # 2. 尋找標題行
            header_row_idx = -1
            user_col_idx = -1
            
            # 關鍵字：填單人員
            target_col = "填單人員"
            
            # 檢查前 20 行
            for idx, row in df_raw.head(20).iterrows():
                # 將 row 轉為 string 並搜尋
                row_values = [str(x).strip() for x in row.values]
                if target_col in row_values:
                    header_row_idx = idx
                    break
            
            if header_row_idx == -1:
                print(f"[WARNING] 未找到包含 '{target_col}' 的標題行，嘗試直接讀取第一行")
                header_row_idx = 0
            
            # 3. 重新讀取，設定 header
            df = pd.read_excel(file_path, header=header_row_idx)
            
            # 清理欄位名稱 (移除前後空白)
            df.columns = df.columns.astype(str).str.strip()
            
            # 再次確認 '填單人員' 欄位是否存在
            if target_col not in df.columns:
                print(f"[ERROR] 找不到 '{target_col}' 欄位")
                # 嘗試模糊搜尋
                possible_cols = [c for c in df.columns if '填單' in str(c) or '人員' in str(c)]
                if possible_cols:
                    print(f"[INFO] 使用近似欄位: {possible_cols[0]}")
                    target_col = possible_cols[0]
                else:
                    return None
            
            # 4. 過濾無效資料 (填單人員為空的行)
            df = df.dropna(subset=[target_col])
            
            # 5. 分組
            users = df[target_col].unique().tolist()
            users = [str(u).strip() for u in users if str(u).strip()]
            users = sorted(list(set(users)))
            
            has_multiple_users = len(users) > 1
            data_by_user = {}
            
            # 用於全量數據的 Markdown
            all_data = df.to_markdown(index=False)
            
            # 6. 為每個人生成資料
            for user in users:
                user_df = df[df[target_col] == user]
                
                # 處理日期格式，避免輸出 Timestamp('...')
                for col in user_df.columns:
                    if pd.api.types.is_datetime64_any_dtype(user_df[col]):
                        user_df[col] = user_df[col].dt.strftime('%Y-%m-%d')
                
                # 轉為 Markdown 表格字串
                user_data_str = user_df.to_markdown(index=False)
                data_by_user[user] = user_data_str
                
            return {
                'has_multiple_users': has_multiple_users,
                'users': users,
                'data_by_user': data_by_user,
                'all_data': all_data,
                'df': df # 保留 DataFrame 以供後續需要
            }
            
        except Exception as e:
            print(f"[ERROR] Pandas 解析失敗: {e}")
            import traceback
            traceback.print_exc()
            return None

    @staticmethod
    def parse_work_report_excel(file_path):
        """解析工作報告 Excel，返回結構化資料
        
        Args:
            file_path: Excel 檔案路徑
            
        Returns:
            dict: {
                'has_multiple_users': bool,
                'users': list,
                'data_by_user': dict,
                'all_data': str
            }
        """
        # 優先嘗試 Pandas
        result = ExcelParser.parse_with_pandas(file_path)
        if result:
            return result
            
        print("[WARNING] Pandas 解析失敗，嘗試使用舊版解析器 (Fallback)")
        
        # 舊版邏輯 fallback
        ext = file_path.lower().split('.')[-1]
        parsed_data = None
        if ext == 'xlsx':
            parsed_data = ExcelParser.parse_xlsx(file_path)
        elif ext == 'xls':
            parsed_data = ExcelParser.parse_xls(file_path)
        
        if parsed_data:
            return ExcelParser.group_by_user(parsed_data)
        
        return {
            'has_multiple_users': False,
            'users': [],
            'data_by_user': {},
            'all_data': ''
        }

    @staticmethod
    def find_user_column(headers):
        """(Legacy) 在標題行中尋找填單人員欄位"""
        for idx, header in enumerate(headers):
            if header:
                header_raw = str(header)
                header_str = header_raw.strip().lower()
                header_clean = header_str.replace(" ", "").replace("　", "")
                
                for keyword in ExcelParser.USER_FIELD_KEYWORDS:
                    if keyword in header_clean:
                        return idx
        return None
    
    @staticmethod
    def parse_xlsx(file_path):
        """(Legacy) 解析 XLSX 檔案"""
        if not OPENPYXL_AVAILABLE:
            return None
        
        try:
            wb = load_workbook(file_path, data_only=True)
            all_records = []
            user_column = None
            headers = []
            
            for sheet_name in wb.sheetnames:
                sheet = wb[sheet_name]
                rows = list(sheet.iter_rows())
                
                if not rows:
                    continue
                
                header_row_index = -1
                for i in range(min(len(rows), 10)):
                    row_values = [str(cell.value) if cell.value else '' for cell in rows[i]]
                    found_col = ExcelParser.find_user_column(row_values)
                    if found_col is not None:
                        headers = row_values
                        user_column = found_col
                        header_row_index = i
                        break
                
                if header_row_index == -1:
                    if rows:
                        headers = [str(cell.value) if cell.value else '' for cell in rows[0]]
                        header_row_index = 0
                
                for row_idx, row in enumerate(rows[header_row_index+1:]):
                    row_values = [str(cell.value) if cell.value else '' for cell in row]
                    if any(row_values):
                        user_val = None
                        if user_column is not None and user_column < len(row_values):
                            user_val = row_values[user_column].strip()
                            
                        record = {
                            'sheet': sheet_name,
                            'data': row_values,
                            'user': user_val
                        }
                        all_records.append(record)
            
            wb.close()
            return {
                'headers': headers,
                'user_column': user_column,
                'records': all_records
            }
        except Exception:
            return None
    
    @staticmethod
    def parse_xls(file_path):
        """(Legacy) 解析 XLS 檔案"""
        if XLRD_AVAILABLE:
            try:
                wb = xlrd.open_workbook(file_path)
                all_records = []
                user_column = None
                headers = []
                
                for sheet_idx in range(wb.nsheets):
                    sheet = wb.sheet_by_index(sheet_idx)
                    if sheet.nrows == 0: continue
                    
                    header_row_index = -1
                    for i in range(min(sheet.nrows, 10)):
                        row_values = [str(sheet.cell_value(i, col_idx)) for col_idx in range(sheet.ncols)]
                        found_col = ExcelParser.find_user_column(row_values)
                        if found_col is not None:
                            headers = row_values
                            user_column = found_col
                            header_row_index = i
                            break
                    
                    if header_row_index == -1:
                        header_row_index = 0
                        headers = [str(sheet.cell_value(0, col_idx)) for col_idx in range(sheet.ncols)]
                    
                    for row_idx in range(header_row_index + 1, sheet.nrows):
                        row_values = [str(sheet.cell_value(row_idx, col_idx)) for col_idx in range(sheet.ncols)]
                        if any(row_values):
                            user_val = None
                            if user_column is not None and user_column < len(row_values):
                                user_val = row_values[user_column].strip()
                            record = {'sheet': sheet.name, 'data': row_values, 'user': user_val}
                            all_records.append(record)
                
                return {'headers': headers, 'user_column': user_column, 'records': all_records}
            except Exception:
                pass
        
        # COM Fallback ... (Skipping full COM impl for brevity unless critical)
        return None

    @staticmethod
    def group_by_user(parsed_data):
        """(Legacy) 將解析後的資料按人員分組"""
        if not parsed_data or not parsed_data['records']:
            return {'has_multiple_users': False, 'users': [], 'data_by_user': {}, 'all_data': ''}
        
        headers = parsed_data['headers']
        records = parsed_data['records']
        
        users_set = set()
        for record in records:
            if record['user'] and record['user'].strip():
                users_set.add(record['user'])
        
        users = sorted(list(users_set))
        has_multiple_users = len(users) > 1
        data_by_user = {}
        
        if has_multiple_users:
            for user in users:
                user_records = [r for r in records if r['user'] == user]
                data_by_user[user] = ExcelParser.format_records_as_text(headers, user_records)
        
        all_data = ExcelParser.format_records_as_text(headers, records)
        return {
            'has_multiple_users': has_multiple_users,
            'users': users,
            'data_by_user': data_by_user,
            'all_data': all_data
        }
    
    @staticmethod
    def format_records_as_text(headers, records):
        """(Legacy) 將記錄格式化為文字"""
        if not records: return ""
        lines = []
        current_sheet = None
        for record in records:
            if record['sheet'] != current_sheet:
                current_sheet = record['sheet']
                lines.append(f"\n=== 工作表: {current_sheet} ===")
                lines.append(" | ".join(headers))
            lines.append(" | ".join(record['data']))
        return "\n".join(lines)
